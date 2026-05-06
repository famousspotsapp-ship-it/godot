#!/usr/bin/env python3
"""Generate comic_translation_asset_checklist.xlsx.

Standalone script that builds a two-sheet workbook:
  1. "Asset Checklist" — categorized asset requirements with priority,
     status dropdown, and fallback steps.
  2. "Process Overview" — three-phase rollout plan.

Usage:
    pip install openpyxl
    python3 generate_checklist.py
"""

from __future__ import annotations

import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------------------
# Data
# ---------------------------------------------------------------------------

HEADERS = [
    "Category",
    "Asset",
    "Format Needed",
    "Why It's Needed",
    "Priority",
    "Has This? (Y/N/Partial)",
    "If No — Fallback Steps",
]

ROWS: list[list[str]] = [
    # Source Files
    [
        "Source Files",
        "Layered art files (text on separate layers)",
        ".psd / .clip / .kra",
        "Allows swapping text without touching artwork — the single most important asset",
        "Critical",
        "",
        "Request textless exports instead (see row below). If neither "
        "available, will need AI inpainting to remove text from final pages.",
    ],
    [
        "Source Files",
        "Textless page exports (art without any text)",
        ".png or .tiff at 300+ DPI",
        "Clean backgrounds behind every bubble — eliminates need for inpainting",
        "Critical (if no layered files)",
        "",
        "Will use AI inpainting (IOPaint/LaMa) to erase text from the final "  # codespell:ignore-line
        "lettered pages. Quality depends on art complexity behind bubbles.",
    ],
    # Fonts
    [
        "Fonts",
        "All font files used in lettering",
        ".ttf or .otf",
        "Exact font matching for the French text — biggest factor in visual quality",
        "Critical",
        "",
        "Will need to visually identify fonts using tools like WhatTheFont or "
        "Identifont, then source matching or similar free/commercial fonts.",
    ],
    [
        "Fonts",
        "Font usage map (which font is used where)",
        "Any text format (.txt, .docx)",
        "Knowing which font applies to dialogue, narration, SFX, titles, character-specific styles",  # codespell:ignore dialogue
        "Important",
        "",
        "Will manually catalog font usage by reviewing the final lettered pages — slower but doable.",
    ],
    # Text & Translation Reference
    [
        "Text & Translation Reference",
        "Full script / dialogue transcript",  # codespell:ignore dialogue
        ".txt, .docx, .csv, or .xlsx",
        "Every line of dialogue keyed to page and panel number — avoids OCR errors and gives translator full context",  # codespell:ignore dialogue
        "Critical",
        "",
        "Will run OCR (Tesseract/EasyOCR) on each bubble to extract text. Will need a manual review pass for accuracy.",
    ],
    [
        "Text & Translation Reference",
        "Glossary of names and invented terms",
        "Any text format",
        "Ensures consistent translation of proper nouns, place names, and made-up words across all pages",
        "Important",
        "",
        "Will build glossary manually during translation by cataloging "
        "recurring names/terms from the script. Creator should review before "
        "final render.",
    ],
    [
        "Text & Translation Reference",
        "Style/tone notes",
        "Any text format",
        "Character voice descriptions, humor notes, cultural references needing adaptation vs. literal translation",
        "Nice-to-have",
        "",
        "Translator will infer tone from context. A review pass with the creator after first draft is recommended.",
    ],
    # Layout Reference
    [
        "Layout Reference",
        "Final lettered PDF or images",
        ".pdf or .png",
        "Reference for text placement, sizing, emphasis (bold/italic), and bubble assignment",
        "Critical",
        "",
        "This is presumably what we already have — the published comic. If not, need at minimum page images.",
    ],
    [
        "Layout Reference",
        "Balloon/bubble templates (if created separately)",
        ".psd or .png",
        "Some letterers create bubbles as separate elements — having these isolated simplifies text replacement",
        "Nice-to-have",
        "",
        "Will detect and preserve bubbles from the final pages using OpenCV contour detection.",
    ],
    # Nice-to-haves
    [
        "Nice-to-haves",
        "SFX (sound effects) as separate layers or list",
        "Layered file or text list",
        "Sound effects baked into art are hardest to replace. Separate layers make it trivial.",
        "Important",
        "",
        "Creator provides a list of which SFX to translate vs. keep as-is. "
        "For baked-in SFX, will use inpainting + re-rendering (quality may "
        "vary).",
    ],
    [
        "Nice-to-haves",
        "Page dimensions and bleed/trim specs",
        "Text description or PDF preset file",
        "Ensures final French PDF matches original print/digital specifications exactly",
        "Nice-to-have",
        "",
        "Will extract dimensions from the source PDF metadata. Creator confirms if correct.",
    ],
]

PHASES: list[tuple[str, list[str]]] = [
    (
        "Phase 1 — Pilot (Single Issue, English → French)",
        [
            "Collect assets using Sheet 1 checklist",
            "Process one issue end-to-end as proof of concept",
            "Creator reviews and approves quality",
            "Document the pipeline and any issue-specific adjustments",
        ],
    ),
    (
        "Phase 2 — Scale Languages (Single Issue, Multiple Languages)",
        [
            "Reuse the same pipeline with different translation targets",
            "One glossary per target language (reviewed by native speakers)",
            "Verify font files support required character sets (e.g., accents for French, umlauts for German, etc.)",
        ],
    ),
    (
        "Phase 3 — Scale Catalog (All Comics, All Languages)",
        [
            "Apply the proven pipeline to all other comics",
            "Batch process with consistent tooling",
            "Per-comic asset collection using the same checklist template",
        ],
    ),
]


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

BODY_ALIGN = Alignment(vertical="top", wrap_text=True)

ALT_ROW_FILL = PatternFill("solid", fgColor="F2F2F2")

CRITICAL_FILL = PatternFill("solid", fgColor="F8CBAD")  # soft red
IMPORTANT_FILL = PatternFill("solid", fgColor="FFE699")  # soft yellow
NICE_FILL = PatternFill("solid", fgColor="C6EFCE")  # soft green

THIN_SIDE = Side(style="thin", color="BFBFBF")
THIN_BORDER = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)

PHASE_TITLE_FONT = Font(bold=True, size=12, color="FFFFFF")
PHASE_TITLE_FILL = PatternFill("solid", fgColor="2E75B6")
PHASE_TITLE_ALIGN = Alignment(horizontal="left", vertical="center", indent=1)

BULLET_ALIGN = Alignment(vertical="top", wrap_text=True, indent=2)

OVERVIEW_HEADER_FONT = Font(bold=True, size=14, color="FFFFFF")
OVERVIEW_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
OVERVIEW_HEADER_ALIGN = Alignment(horizontal="center", vertical="center")


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------


def _autosize_columns(ws, min_widths: dict[int, int] | None = None, max_width: int = 60) -> None:
    """Set column widths based on the longest line in each cell.

    `min_widths` maps 1-based column index to a minimum width.
    """
    min_widths = min_widths or {}
    for col_idx, column_cells in enumerate(ws.columns, start=1):
        longest = 0
        for cell in column_cells:
            value = cell.value
            if value is None:
                continue
            for line in str(value).splitlines() or [str(value)]:
                if len(line) > longest:
                    longest = len(line)
        width = min(max(longest + 2, min_widths.get(col_idx, 10)), max_width)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def build_asset_checklist(ws) -> None:
    ws.title = "Asset Checklist"

    # Header row
    ws.append(HEADERS)
    for col_idx, _ in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[1].height = 32

    # Data rows
    for row_offset, row_data in enumerate(ROWS, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_offset, column=col_idx, value=value)
            cell.alignment = BODY_ALIGN
            cell.border = THIN_BORDER
            # Alternating row fill (skip Priority cells so CF can win there).
            if row_offset % 2 == 0 and col_idx != 5:
                cell.fill = ALT_ROW_FILL

    last_row = 1 + len(ROWS)
    last_col_letter = get_column_letter(len(HEADERS))
    data_range = f"A2:{last_col_letter}{last_row}"

    # Freeze header row
    ws.freeze_panes = "A2"

    # Auto filter on the full data range (incl. header)
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"

    # Column widths
    _autosize_columns(
        ws,
        min_widths={
            1: 22,  # Category
            2: 32,  # Asset
            3: 22,  # Format Needed
            4: 40,  # Why It's Needed
            5: 22,  # Priority
            6: 18,  # Has This?
            7: 50,  # Fallback Steps
        },
        max_width=55,
    )

    # Reasonable row heights for wrapped text
    for r in range(2, last_row + 1):
        ws.row_dimensions[r].height = 60

    # Data validation dropdown for "Has This?" column (F)
    dv = DataValidation(
        type="list",
        formula1='"Yes,No,Partial"',
        allow_blank=True,
        showDropDown=False,  # False = the dropdown arrow IS shown
    )
    dv.error = "Please choose Yes, No, or Partial."
    dv.errorTitle = "Invalid value"
    dv.prompt = "Yes / No / Partial"
    dv.promptTitle = "Has This?"
    dv.add(f"F2:F{last_row}")
    ws.add_data_validation(dv)

    # Conditional formatting for Priority column (E)
    priority_range = f"E2:E{last_row}"
    ws.conditional_formatting.add(
        priority_range,
        FormulaRule(
            formula=['ISNUMBER(SEARCH("Critical",E2))'],
            stopIfTrue=True,
            fill=CRITICAL_FILL,
        ),
    )
    ws.conditional_formatting.add(
        priority_range,
        FormulaRule(
            formula=['ISNUMBER(SEARCH("Important",E2))'],
            stopIfTrue=True,
            fill=IMPORTANT_FILL,
        ),
    )
    ws.conditional_formatting.add(
        priority_range,
        FormulaRule(
            formula=['ISNUMBER(SEARCH("Nice-to-have",E2))'],
            stopIfTrue=True,
            fill=NICE_FILL,
        ),
    )

    # Apply borders to the whole data range (already applied per-cell above,
    # but be explicit for any cells we may have missed).
    for row in ws[data_range]:
        for cell in row:
            cell.border = THIN_BORDER


def build_process_overview(ws) -> None:
    ws.title = "Process Overview"

    # Title row spanning the layout
    ws.merge_cells("A1:B1")
    title = ws.cell(row=1, column=1, value="Comic Translation Pipeline — Process Overview")
    title.font = OVERVIEW_HEADER_FONT
    title.fill = OVERVIEW_HEADER_FILL
    title.alignment = OVERVIEW_HEADER_ALIGN
    ws.row_dimensions[1].height = 30

    current_row = 3

    for phase_title, bullets in PHASES:
        # Phase title spans columns A-B
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
        cell = ws.cell(row=current_row, column=1, value=phase_title)
        cell.font = PHASE_TITLE_FONT
        cell.fill = PHASE_TITLE_FILL
        cell.alignment = PHASE_TITLE_ALIGN
        cell.border = THIN_BORDER
        # Border on merged-away cell so the bottom edge renders properly
        ws.cell(row=current_row, column=2).border = THIN_BORDER
        ws.row_dimensions[current_row].height = 24
        current_row += 1

        for bullet in bullets:
            marker = ws.cell(row=current_row, column=1, value="•")
            marker.alignment = Alignment(horizontal="center", vertical="top")
            marker.border = THIN_BORDER
            text = ws.cell(row=current_row, column=2, value=bullet)
            text.alignment = BULLET_ALIGN
            text.border = THIN_BORDER
            ws.row_dimensions[current_row].height = 22
            current_row += 1

        # Spacer row between phases
        current_row += 1

    # Column widths
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 95

    # Freeze the title row
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main() -> None:
    wb = Workbook()
    build_asset_checklist(wb.active)
    build_process_overview(wb.create_sheet("Process Overview"))

    output_path = Path.cwd() / "comic_translation_asset_checklist.xlsx"
    wb.save(output_path)
    print(os.path.abspath(output_path))


if __name__ == "__main__":
    main()
